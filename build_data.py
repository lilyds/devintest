#!/usr/bin/env python3
"""Build ae_data.json for the AE Targets dashboard from the consumption workbook.
All figures derived from real weekly consumption data. April daily shape is simulated
in the frontend from each account's 4 April weekly buckets (full-month totals stay exact)."""
import openpyxl, json, random
from collections import defaultdict

CONS = '/home/ubuntu/attachments/d956713d-4790-4dc1-a6de-8c424002a57e/GTM+Finance+Consumption+Data_Lily.xlsx'

STRETCH = {'Demo':1.20,'POC':1.15,'UAT':1.15,'Production':1.08,'Internal':1.00,
           'Archived':0.0,'Unknown':1.05}
MARCH_WEEKS = ['2026-03-02','2026-03-09','2026-03-16','2026-03-23','2026-03-30']
APRIL_WEEKS = ['2026-04-06','2026-04-13','2026-04-20','2026-04-27']
# full ordered week list for forecasting (anchor = last March week 2026-03-30)
ALL_WEEKS = ['2025-11-03','2025-11-10','2025-11-17','2025-11-24','2025-12-01','2025-12-08',
 '2025-12-15','2025-12-22','2025-12-29','2026-01-05','2026-01-12','2026-01-19','2026-01-26',
 '2026-02-02','2026-02-09','2026-02-16','2026-02-23','2026-03-02','2026-03-09','2026-03-16',
 '2026-03-23','2026-03-30']  # through end of March (anchor weeks for the April forecast)

# EWA forecast params (mirror model Assumptions)
ALPHA=0.6; LOOKBACK=6; CLIP_MAX=0.12; CLIP_MIN=-0.15; DECAY=0.85

wb = openpyxl.load_workbook(CONS, data_only=True, read_only=True)

# --- price map from Portfolio Data ---
price_map={}
pd=wb['Portfolio Data']
for r in pd.iter_rows(min_row=2, values_only=True):
    eid=r[0]; eff=r[2]
    if isinstance(eid,str) and isinstance(eff,(int,float)) and eff>0:
        price_map[eid]=float(eff)

# --- consumption aggregation ---
ws=wb['Cognition Consumption Data']
it=ws.iter_rows(values_only=True); next(it)
ent_week=defaultdict(lambda: defaultdict(float))   # eid -> week -> acus
ent_phase_count=defaultdict(lambda: defaultdict(float))  # eid -> phase -> acus (weight by acus)
ent_price_num=defaultdict(float); ent_price_den=defaultdict(float)  # acu-weighted price fallback
all_price_num=0.0; all_price_den=0.0
for row in it:
    acc,eid,sf,phase,plan,prod,wk,price,acus,be,au,ao=row
    if not eid: continue
    w=str(wk)[:10]
    a=float(acus or 0)
    ent_week[eid][w]+=a
    if phase: ent_phase_count[eid][phase]+=max(a,0.0)+0.001
    if price not in (None,'') and a>0:
        ent_price_num[eid]+=float(price)*a; ent_price_den[eid]+=a
        all_price_num+=float(price)*a; all_price_den+=a
GLOBAL_PRICE=round(all_price_num/all_price_den,2) if all_price_den else 2.5

def ewa(vals):
    e=vals[0]
    for v in vals[1:]: e=ALPHA*v+(1-ALPHA)*e
    return e

def april_forecast(weekvals):
    """weekvals: dict week->acu through March. Returns 4 forecast weekly ACUs for April."""
    series=[weekvals.get(w,0.0) for w in ALL_WEEKS]
    # anchor = last nonzero week
    anchor=0.0
    for v in reversed(series):
        if v>0: anchor=v; break
    if anchor<=0: return [0.0,0.0,0.0,0.0]
    # WoW growth over recent LOOKBACK weeks (where prev>0)
    growths=[]
    for i in range(len(series)-LOOKBACK, len(series)):
        if i<=0: continue
        prev=series[i-1]; cur=series[i]
        if prev>0: growths.append(cur/prev-1.0)
    g=ewa(growths) if growths else 0.0
    g=max(CLIP_MIN, min(CLIP_MAX, g))
    out=[]; level=anchor; gg=g
    for k in range(4):
        gg = g*(DECAY**k)
        level=level*(1+gg)
        out.append(level)
    return out

accounts=[]
for eid, weekvals in ent_week.items():
    march=sum(weekvals.get(w,0.0) for w in MARCH_WEEKS)
    april_weeks=[round(weekvals.get(w,0.0),2) for w in APRIL_WEEKS]
    # phase = acu-weighted dominant phase (fallback Unknown)
    pc=ent_phase_count.get(eid,{})
    phase=max(pc, key=pc.get) if pc else 'Unknown'
    if phase not in STRETCH: phase='Unknown'
    # price
    if eid in price_map: price=round(price_map[eid],2)
    elif ent_price_den.get(eid,0)>0: price=round(ent_price_num[eid]/ent_price_den[eid],2)
    else: price=GLOBAL_PRICE
    fc=april_forecast(weekvals)
    target=round(sum(fc)*STRETCH[phase],2)
    accounts.append(dict(id=eid, phase=phase, march=round(march,2),
                         weeks=april_weeks, target=target, price=price))

# --- deterministic random AE assignment (all accounts) ---
accounts.sort(key=lambda x:x['id'])
rng=random.Random(42)
rng.shuffle(accounts)
NAE=10
for i,a in enumerate(accounts):
    a['ae']='AE'+str(i%NAE+1)
accounts.sort(key=lambda x:x['id'])

meta=dict(month='April 2026', prev_month='March 2026', days_in_month=30,
          default_asof=5, april_weeks=APRIL_WEEKS, stretch=STRETCH,
          global_price=GLOBAL_PRICE, n_accounts=len(accounts))
json.dump(dict(meta=meta, accounts=accounts),
          open('/home/ubuntu/repos/devintest/ae_data.json','w'))

# sanity prints
tot_march=sum(a['march'] for a in accounts)
tot_april=sum(sum(a['weeks']) for a in accounts)
tot_target=sum(a['target'] for a in accounts)
print('accounts',len(accounts),'global_price',GLOBAL_PRICE)
print('March total',round(tot_march),'(model 16,670,304)')
print('April total',round(tot_april),'(model 18,171,836)')
print('Target total',round(tot_target))
from collections import Counter
print('per-AE counts',dict(Counter(a['ae'] for a in accounts)))
print('phase dist',dict(Counter(a['phase'] for a in accounts)))
